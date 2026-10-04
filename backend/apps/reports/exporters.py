from __future__ import annotations

import csv
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path

from django.conf import settings
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table as ExcelTable
from openpyxl.worksheet.table import TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    HRFlowable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)
from reportlab.platypus import (
    Image as PdfImage,
)
from reportlab.platypus import (
    Table as PdfTable,
)
from reportlab.platypus import (
    TableStyle as PdfTableStyle,
)

BRAND = "009688"
BRAND_DARK = "075E54"
INK = "10233F"
MUTED = "64748B"
LIGHT = "EAF8F6"
GRID = "DDE4EA"


def pdf_color(value):
    return colors.HexColor(f"#{value}")


def logo_path(*, dark: bool = False) -> Path | None:
    filename = "situr-smart-logo-dark-v2.png"
    candidates = (
        Path(__file__).resolve().parent / "assets" / filename,
        Path(settings.BASE_DIR).parent / "web" / "public" / "branding" / filename,
        Path(settings.BASE_DIR) / "static" / "branding" / filename,
    )
    return next((path for path in candidates if path.is_file()), None)


def safe_csv_value(value):
    text = "" if value is None else str(value)
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def _parse_datetime(value):
    if not isinstance(value, str):
        return None
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return None
    if timezone.is_aware(parsed):
        parsed = timezone.localtime(parsed)
    return parsed


def export_value(key, value, *, excel=False):
    if key not in {"registro", "fecha"}:
        return value
    parsed = _parse_datetime(value)
    if not parsed:
        return value
    if excel:
        return parsed.replace(tzinfo=None)
    return parsed.strftime("%d/%m/%Y %H:%M")


def report_title(report):
    return {
        "plataforma": "Empresas y planes",
        "catalogo": "Catálogo turístico",
        "hospedajes": "Hospedajes y habitaciones",
        "actividad": "Actividad de la plataforma",
    }.get(report["tipo"], "Reporte")


def generated_at():
    return timezone.localtime().strftime("%d/%m/%Y %H:%M")


def filter_labels(report):
    filters = report.get("filtros", {})
    date_from = filters.get("desde") or "Inicio"
    date_to = filters.get("hasta") or "Actualidad"
    return {
        "alcance": report["empresa"]["nombre"] if report.get("empresa") else "Toda la plataforma",
        "periodo": f"{date_from} a {date_to}",
        "filtro": filters.get("estado") or "Todos los estados",
        "registros": len(report["filas"]),
    }


def build_csv(report) -> str:
    output = StringIO()
    output.write("\ufeff")
    writer = csv.writer(output)
    writer.writerow(["SITUR-SMART", report_title(report)])
    writer.writerow(["Alcance", report["empresa"]["nombre"] if report.get("empresa") else "Toda la plataforma"])
    writer.writerow(["Generado", generated_at()])
    filters = filter_labels(report)
    writer.writerow(["Período", filters["periodo"]])
    writer.writerow(["Filtro", filters["filtro"]])
    writer.writerow([])
    writer.writerow(["Indicador", "Valor", "Detalle"])
    for metric in report["indicadores"]:
        writer.writerow([metric["etiqueta"], metric["valor"], metric.get("detalle", "")])
    writer.writerow([])
    writer.writerow([column[1] for column in report["columnas"]])
    for row in report["filas"]:
        writer.writerow([
            safe_csv_value(export_value(column[0], row.get(column[0])))
            for column in report["columnas"]
        ])
    writer.writerow([])
    writer.writerow(["Nota", report["nota"]])
    return output.getvalue()


def build_xlsx(report) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reporte"
    columns = report["columnas"]
    width = max(len(columns), 6)

    logo = logo_path(dark=True)
    title_column = 4 if logo else 1
    if logo:
        sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=3)
        sheet.cell(1, 1).fill = PatternFill("solid", fgColor=BRAND_DARK)
        image = ExcelImage(logo)
        image.width = 205
        image.height = 68
        sheet.add_image(image, "A1")
    sheet.merge_cells(start_row=1, start_column=title_column, end_row=1, end_column=width)
    title = sheet.cell(1, title_column, report_title(report))
    title.font = Font(name="Aptos Display", size=20, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor=BRAND_DARK)
    title.alignment = Alignment(vertical="center", horizontal="left")
    sheet.row_dimensions[1].height = 38

    filters = filter_labels(report)
    scope = filters["alcance"]
    sheet.merge_cells(start_row=2, start_column=title_column, end_row=2, end_column=width)
    meta = sheet.cell(2, title_column, f"Alcance: {scope}   •   Generado: {generated_at()}")
    meta.font = Font(size=10, color="D6EFEC")
    meta.fill = PatternFill("solid", fgColor=BRAND_DARK)
    meta.alignment = Alignment(vertical="top")
    sheet.row_dimensions[2].height = 24

    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=width)
    summary = sheet.cell(
        3, 1,
        f"ALCANCE: {scope}     PERÍODO: {filters['periodo']}     FILTRO: {filters['filtro']}     REGISTROS: {filters['registros']}",
    )
    summary.font = Font(size=9, bold=True, color=INK)
    summary.fill = PatternFill("solid", fgColor=LIGHT)
    summary.alignment = Alignment(vertical="center")
    sheet.row_dimensions[3].height = 24

    for index, metric in enumerate(report["indicadores"], start=1):
        if index > width:
            break
        cell = sheet.cell(5, index, metric["etiqueta"])
        cell.font = Font(size=9, bold=True, color=MUTED)
        cell.fill = PatternFill("solid", fgColor=LIGHT)
        value = sheet.cell(6, index, metric["valor"])
        value.font = Font(size=16, bold=True, color=INK)
        value.fill = PatternFill("solid", fgColor=LIGHT)
        value.alignment = Alignment(vertical="top")
    sheet.row_dimensions[6].height = 28

    header_row = 8
    for col, (_, label) in enumerate(columns, start=1):
        cell = sheet.cell(header_row, col, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BRAND)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 28

    for row_index, row in enumerate(report["filas"], start=header_row + 1):
        for col_index, (key, _) in enumerate(columns, start=1):
            cell = sheet.cell(row_index, col_index, export_value(key, row.get(key, ""), excel=True))
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if key in {"registro", "fecha"} and isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F5FAFA")

    if report["filas"]:
        table_ref = f"A{header_row}:{get_column_letter(len(columns))}{header_row + len(report['filas'])}"
        table = ExcelTable(displayName="DatosReporte", ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)

    note_row = header_row + max(len(report["filas"]), 1) + 2
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=width)
    note = sheet.cell(note_row, 1, "Nota: " + report["nota"])
    note.font = Font(size=9, italic=True, color=MUTED)
    note.alignment = Alignment(wrap_text=True)

    thin = Side(style="thin", color=GRID)
    for row in sheet.iter_rows(min_row=7, max_row=header_row + len(report["filas"]), max_col=len(columns)):
        for cell in row:
            cell.border = Border(bottom=thin)

    for index, (key, label) in enumerate(columns, start=1):
        values = [str(export_value(key, row.get(key, ""))) for row in report["filas"][:200]]
        longest = max([len(label), *(len(value) for value in values)], default=len(label))
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 12), 34)

    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{header_row + max(len(report['filas']), 1)}"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"1:{header_row}"
    sheet.oddFooter.center.text = "SITUR-SMART"
    sheet.oddFooter.right.text = "Página &P de &N"
    sheet.oddFooter.left.text = "&D &T"

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _pdf_value(value):
    text = "—" if value is None or value == "" else str(value)
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_pdf(report) -> bytes:
    stream = BytesIO()
    page_size = landscape(A4)
    document = SimpleDocTemplate(
        stream, pagesize=page_size, leftMargin=14 * mm, rightMargin=14 * mm,
        topMargin=18 * mm, bottomMargin=16 * mm, title=report_title(report), author="SITUR-SMART",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Brand", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=20, leading=23, textColor=pdf_color(INK), spaceAfter=2))
    styles.add(ParagraphStyle(name="Meta", parent=styles["Normal"], fontSize=8.5, leading=12, textColor=pdf_color(MUTED)))
    styles.add(ParagraphStyle(name="MetricLabel", parent=styles["Normal"], fontSize=7.5, leading=10, textColor=pdf_color(MUTED)))
    styles.add(ParagraphStyle(name="MetricValue", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=pdf_color(INK)))
    styles.add(ParagraphStyle(name="Cell", parent=styles["Normal"], fontSize=7, leading=9, textColor=pdf_color(INK)))
    styles.add(ParagraphStyle(name="CellHead", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=7, leading=9, textColor=colors.white))

    filters = filter_labels(report)
    scope = filters["alcance"]
    control = PdfTable(
        [[Paragraph(generated_at(), styles["Meta"]), Paragraph("Reporte operativo - SITUR-SMART", styles["Meta"]) ]],
        colWidths=[(page_size[0] - 28 * mm) / 2] * 2,
    )
    control.setStyle(PdfTableStyle([
        ("ALIGN", (1, 0), (1, 0), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    report_heading = [
        Paragraph(report_title(report).upper(), ParagraphStyle(name="ReportTitle", parent=styles["Heading2"], fontSize=13, leading=16, alignment=2, textColor=pdf_color(BRAND), spaceAfter=3)),
        Paragraph(f"Emisión: {generated_at()}", ParagraphStyle(name="Emission", parent=styles["Meta"], alignment=2)),
    ]
    logo = logo_path()
    if logo:
        logo_image = PdfImage(str(logo), width=55 * mm, height=55 * mm / 3)
        header = PdfTable([[logo_image, report_heading]], colWidths=[62 * mm, page_size[0] - 90 * mm])
        header.setStyle(PdfTableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BACKGROUND", (0, 0), (0, 0), pdf_color(BRAND_DARK)),
            ("LEFTPADDING", (0, 0), (0, 0), 4),
            ("RIGHTPADDING", (0, 0), (0, 0), 4),
            ("TOPPADDING", (0, 0), (0, 0), 3),
            ("BOTTOMPADDING", (0, 0), (0, 0), 3),
            ("LEFTPADDING", (1, 0), (1, 0), 0),
            ("RIGHTPADDING", (1, 0), (1, 0), 0),
            ("TOPPADDING", (1, 0), (1, 0), 0),
            ("BOTTOMPADDING", (1, 0), (1, 0), 0),
        ]))
    else:
        header = PdfTable([[Paragraph("SITUR-<font color='#009688'>SMART</font>", styles["Brand"]), report_heading]], colWidths=[62 * mm, page_size[0] - 90 * mm])
    filter_data = [[
        Paragraph(f"<b>Alcance:</b> {_pdf_value(scope)}", styles["Meta"]),
        Paragraph(f"<b>Período:</b> {_pdf_value(filters['periodo'])}", styles["Meta"]),
        Paragraph(f"<b>Filtro:</b> {_pdf_value(filters['filtro'])}", styles["Meta"]),
        Paragraph(f"<b>Registros:</b> {filters['registros']}", ParagraphStyle(name="RecordCount", parent=styles["Meta"], alignment=2)),
    ]]
    filter_table = PdfTable(filter_data, colWidths=[60 * mm, 58 * mm, 62 * mm, page_size[0] - 208 * mm])
    filter_table.setStyle(PdfTableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), pdf_color("F8FAFC")),
        ("BOX", (0, 0), (-1, -1), 0.5, pdf_color(GRID)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story = [
        control, Spacer(1, 4 * mm), header, Spacer(1, 3 * mm),
        HRFlowable(width="100%", thickness=1.5, color=pdf_color(BRAND), spaceBefore=0, spaceAfter=4 * mm),
        filter_table, Spacer(1, 5 * mm),
    ]

    metric_cells = []
    for metric in report["indicadores"]:
        detail = f" {metric.get('detalle', '')}" if metric.get("detalle") else ""
        metric_cells.append([
            Paragraph(_pdf_value(metric["etiqueta"]), styles["MetricLabel"]),
            Paragraph(f"{_pdf_value(metric['valor'])}{detail}", styles["MetricValue"]),
        ])
    cards = [
        PdfTable(
            [[cell[0]], [cell[1]]],
            colWidths=[(page_size[0] - 28 * mm) / max(len(metric_cells), 1) - 16],
        )
        for cell in metric_cells
    ]
    metric_table = PdfTable([cards], colWidths=[(page_size[0] - 28 * mm) / max(len(cards), 1)] * len(cards))
    metric_table.setStyle(PdfTableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), pdf_color(LIGHT)),
        ("BOX", (0, 0), (-1, -1), 0.5, pdf_color(GRID)),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, pdf_color(GRID)),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend([metric_table, Spacer(1, 6 * mm)])

    columns = report["columnas"]
    data = [[Paragraph(_pdf_value(label), styles["CellHead"]) for _, label in columns]]
    for row in report["filas"]:
        data.append([
            Paragraph(_pdf_value(export_value(key, row.get(key))), styles["Cell"])
            for key, _ in columns
        ])
    if len(data) == 1:
        data.append([Paragraph("Sin datos para los filtros seleccionados", styles["Cell"])] + [""] * (len(columns) - 1))
    available = page_size[0] - 28 * mm
    col_widths = [available / len(columns)] * len(columns)
    table = PdfTable(data, colWidths=col_widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(PdfTableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), pdf_color(BRAND)),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, pdf_color("F5FAFA")]),
        ("GRID", (0, 0), (-1, -1), 0.35, pdf_color(GRID)),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([table, Spacer(1, 4 * mm), Paragraph("Nota: " + _pdf_value(report["nota"]), styles["Meta"])])

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(pdf_color(GRID))
        canvas.line(14 * mm, 11 * mm, page_size[0] - 14 * mm, 11 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(pdf_color(MUTED))
        canvas.drawString(14 * mm, 7 * mm, "SITUR-SMART · Reporte operativo")
        canvas.drawRightString(page_size[0] - 14 * mm, 7 * mm, f"Página {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()
